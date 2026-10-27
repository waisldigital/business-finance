"""Financial Reporting Package → the plan year's P&L actuals (a small workbook in the package's layout):
every block lands in the reports, the Full P&L ties to the package by segment, double-click lists projects,
YTD / MTD filter the older formats, and the package wins over SAP-dump actuals for its months."""
import io
from datetime import datetime

import openpyxl

from aop.package_import import prefer_package

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def fy_months(fy):
    y = 2000 + int(fy[2:]) - 1
    return [datetime(y + (m < 4), m, 1) for m in (4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3)]


def block(ws, title, first_col, label_cols, rows, months, act=2):
    """title row, flags row (Act for the first `act` months, else Fct), dates row, data rows (labels + Cr values)."""
    ws.append([title])
    pad = [None] * (first_col - 1)
    ws.append(pad + [("Act" if i < act else "Fct") for i in range(12)])
    ws.append([None] * (first_col - 1 - len(label_cols)) + list(label_cols) + months)
    for labels, vals in rows:
        ws.append([None] * (first_col - 1 - len(labels)) + list(labels) + list(vals) + [0] * (12 - len(vals)))


def package(fy):
    m = fy_months(fy)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Revenue Analysis"
    block(ws, "SECTION A — CUTE REVENUE ANALYSIS", 4, ["Location", "PAX Type"],
          [(["DIAL", "Domestic"], [20, 22]), (["DIAL", "International"], [10, 11]), (["DIAL", "Combined"], [30, 33]),
           (["GGIAL", "Combined"], [2, 3]), (["TOTAL", "Combined"], [32, 36])], m)
    ws.append([])
    block(ws, "SECTION E — Revenue Share (₹ Crores)", 4, ["Revenue Share", "Location"],
          [(["Rev Share", "DIAL"], [3, 3.5]), (["Rev Share", "Total"], [3, 3.5])], m)
    ws = wb.create_sheet("Non Cute Full Break up")
    block(ws, "x", 5, ["Revenu Type", "Location"], [(["Non - CUTE", "DIAL"], [5, 4]), (["Non - CUTE", "Total"], [5, 4])], m)
    ws = wb.create_sheet("CR Breakup")
    block(ws, "SECTION A — CR", 6, ["Key", "", "", "Location"],
          [(["Change requestDIAL", "", "", "DIAL"], [2, 1]), (["Change requestTotal", "", "", "Total"], [2, 1])], m)
    ws = wb.create_sheet("Project Summary")
    ws.append(["Revenue Analysis — Projects"])
    ws.append([])
    ws.append([])
    ws.append([None] * 9 + ["Revenue (FY27)"] + [None] * 11 + ["Opex (Act FY27)"] + [None] * 11 + ["Direct Resource (Act FY27)"])
    ws.append([None] * 9 + (["Act"] * 2 + ["Fct"] * 10) * 3)
    ws.append(["Key", "WBS", "Project ID", "Project Name", "Segment", "Geography/Region", "Category", "GMR/", "India/"] + m * 3)
    ws.append(["SolutionsSAC", "WSIN.1", "WDP014", "SAC ICT", "Growth", "India", "Projects", "GMR", "India"] + [4, 6] + [0] * 10 + [3, 3] + [0] * 10 + [1, 1] + [0] * 10)
    ws.append(["SolutionsIAG", "WSUK.1", "WDP001", "IAG", "Growth", "EU & UK", "Projects", "Non-GMR", "International"] + [1, 0] + [0] * 10 + [0.5, 0.5] + [0] * 10 + [0, 0.2] + [0] * 10)
    ws = wb.create_sheet("Opex & Dir Resource Airportwsie")
    ws.append(["SECTION A — Opex Cost- Location wise"])
    ws.append([None, None, None] + ["Act", "Act"] + ["Fct"] * 10)
    ws.append([None, "Section", "Airport"] + m)
    ws.append([None, "Opex", "DIAL", 99, 99])  # incl. CR — ignored
    ws.append([None, "Airport Resource Cost- Direct"])
    ws.append([None, "Direct Resource", "DIAL", 2.5, 2.5])
    ws.append([None, "Direct Resource", "Shared services", 0.6, 0.6])
    ws.append([None, "Direct Resource", "Total", 2.5, 2.5])
    ws.append([None, "Shared Service Allocation %"])
    ws.append([None, "DIAL", 1.0])
    ws.append([None, "Without CR"] + [None] + ["Act", "Act"] + ["Fct"] * 10)
    ws.append([None, "Section", "Airport"] + m)
    ws.append([None, "Opex", "DIAL", 5, 6])
    ws.append([None, "Opex", "Shared services", 1, 1])
    ws.append([None, "CR Only"])
    ws.append([None, "Section", "Airport"] + m)
    ws.append([None, "Opex", "DIAL", 0.5, 0.5])
    ws = wb.create_sheet("Resource Cost Raw")
    ws.append([])
    ws.append([])
    ws.append([])
    ws.append([])
    ws.append([None] * 17 + m + [None] + m)
    ws.append(["WBS Code", "For Actuals", "For Budget", "Total Direct"])
    ws.append([None, None, None, "Operations — CA"])
    ws.append(["W1", "DIAL Ops", "DIAL", "DIAL — Ops"] + [None] * 13 + [2, 2] + [0] * 10 + [None] + [100, 101] + [0] * 10)
    ws.append([None, None, None, "Total Indirect"])
    ws.append([None, None, None, "Enabling Functions"])
    ws.append([None, "HR", "HR", "HR"] + [None] * 13 + [0.4, 0.5] + [0] * 10 + [None] + [18, 18] + [0] * 10)
    ws.append([None, None, None, "Pursuit-Oriented Functions"])
    ws.append([None, "Front End Sales - India", "x", "Front-End Sales"] + [None] * 13 + [0.3, 0.3] + [0] * 10 + [None] + [10, 10] + [0] * 10)
    ws.append([None, None, None, "Total Capitalised"])
    ws.append([None, "AODB", "AODB", "AODB"] + [None] * 13 + [9, 9] + [0] * 10)
    ws = wb.create_sheet("OH Summary_Dept")
    block(ws, "SECTION A — Enabiling OH (CA+CR)", 6, ["Key", "Section", "Department"],
          [(["CA+CRAdmin", "CA+CR", "Admin"], [1, 1.2]), (["x", "Total", ""], [1, 1.2])], m)
    block(ws, "Solutions", 6, ["Key", "Section", "Department"],
          [(["SolutionsBD", "Solutions", "BD - International"], [0.9, 1.1])], m)
    ws = wb.create_sheet("P&L")
    ws.append(["WAISL Limited"])
    ws.append(["Monthly MIS"])
    ws.append([None, "Particulars"] + m[:3] + ["Q1"] + m[3:])
    for lab, vals in (("Revenue from Operations", [44, 47]), ("Depreciation", [11, 11.5]), ("Interest & Finance cost", [4, 4]),
                      ("Interest & Finance cost- Prior Period", [0, 1]), ("Other Income", [2, 1]), ("Tax", [2.5, 3])):
        ws.append([None, lab] + vals)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_package_actuals_flow(client, admin):
    cfg = client.get("/api/aop/config", headers=admin).json()
    fy = cfg["plan_fy"]
    months = [d.strftime("%Y-%m") for d in fy_months(fy)]
    r = client.post("/api/aop/import/reporting-package", headers=admin, files={"file": ("pkg.xlsx", package(fy), XLSX)})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["periods"] == months[:2] and out["meta"]["cutoff"] == months[1]
    b = out["meta"]["by_block"]
    assert b["CUTE revenue"] == 63 + 5 and b["Non-CUTE revenue"] == 9 and b["Change request revenue"] == 3
    assert b["Project revenue"] == 11 and b["Opex — Projects"] == 7 and b["Employee cost — direct Projects"] == 2.2
    # airport rows carry their allocated share of shared services: booked net, shared services on their own
    assert b["Opex — CA"] == 11 and b["Opex — CR"] == 1 and b["Employee cost — direct CA+CR"] == 5.0
    assert b["Employee cost — indirect"] == 1.5 and b["Overheads (SG&A)"] == 4.2 and round(b["Below EBITDA"], 2) == 40.0
    assert not out["meta"]["mismatches"]  # ties to the package's own P&L sheet
    assert client.get("/api/aop/config", headers=admin).json()["cutoffs"]["default"] >= months[1]

    pnl = client.get("/api/aop/mis/full-pnl", headers=admin).json()
    rows = {x["id"]: x for x in pnl["rows"]}
    ytd = lambda rid, seg: sum(rows[rid]["values"]["af_plan"][seg][:2]) / 1e7
    assert round(ytd("rev", "cacr"), 3) == 63 + 5 + 9 + 3 - 0 and round(ytd("rev", "sol"), 3) == 11
    assert round(ytd("opex_tp", "cacr"), 3) == 12 and round(ytd("opex_tp", "sol"), 3) == 7
    # indirect costs follow the segment they were booked to (HR → CA+CR, front-end sales → Solutions)
    assert round(ytd("emp_indirect", "cacr"), 3) == 0.9 and round(ytd("emp_indirect", "sol"), 3) == 0.6
    assert round(ytd("sga", "cacr"), 3) == 2.2 and round(ytd("sga", "sol"), 3) == 2.0

    d = client.get("/api/aop/mis/drill", headers=admin, params={"row": "rev_projects"}).json()
    items = {i["label"]: i for i in d["groups"][0]["items"]}
    assert set(items) >= {"SAC ICT", "IAG"} and sum(items["SAC ICT"]["values"]["af_plan"][:2]) == 10e7
    d = client.get("/api/aop/mis/drill", headers=admin, params={"row": "opex_tp"}).json()
    assert [g["label"] for g in d["groups"]][:1] == ["Opex — CA"] and any(g["label"] == "Projects" for g in d["groups"])
    assert client.get("/api/aop/mis/drill", headers=admin, params={"row": "gm"}).json()["groups"] == []

    # older formats take FY / YTD / MTD
    fy_r = client.get("/api/aop/reports/margin", headers=admin, params={"block": "af_plan"}).json()
    mtd = client.get("/api/aop/reports/margin", headers=admin, params={"block": "af_plan", "period": "mtd", "month": 1}).json()
    ytd_r = client.get("/api/aop/reports/margin", headers=admin, params={"block": "af_plan", "period": "ytd", "month": 1}).json()
    tot = lambda x: next(r for r in x["rows"] if r["tag"] == "Total")["revenue"]
    assert mtd["months"] == [months[1]] and len(ytd_r["months"]) == 2 and len(fy_r["months"]) == 12
    assert tot(ytd_r) >= tot(mtd) and tot(mtd) >= 47e7
    for url in ("/api/aop/reports/opex", "/api/aop/reports/overheads", "/api/aop/reports/wbs"):
        assert client.get(url, headers=admin, params={"period": "ytd", "month": 2}).status_code == 200
    fm = client.get("/api/aop/mis/formats", headers=admin).json()
    assert fm["months"]["labels"][0] == "Apr"


def test_package_wins_for_its_months():
    acts = [{"domain": "opex", "period": "2026-04", "amount": 5, "source": "mis"},
            {"domain": "opex", "period": "2026-04", "amount": 7, "source": "package"},
            {"domain": "opex", "period": "2026-05", "amount": 3, "source": "mis"},
            {"domain": "payroll", "period": "2026-04", "amount": 2, "source": "resource"}]
    kept = prefer_package(acts)
    assert [a["amount"] for a in kept] == [7, 3, 2]
    assert prefer_package(acts[:1]) == acts[:1]

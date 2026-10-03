"""End to end through the API (in-memory Mongo): one-time workbook import, ZMM runs, Review (To map, PO changes,
Corrections, Checks, Upload log), drawers, add-on lines and the download → upload round trip (tests 10–11, 13)."""
import io
from datetime import datetime

import openpyxl

from aop import po as P

ZH = P.REQUIRED + ["Currency", "GR Amount In LC", "Nature (Opex/Capex/OH)"]


def zmm_row(po, item, **kw):
    base = {"Purchase Order": po, "Purchase Order Item": item, "Created On": datetime(2026, 4, 1), "Supplier": 1000001,
            "Supplier Name": "Vendor A", "Material": 9700000001, "Material Description": "AMC", "Currency": "INR",
            "Net Order Value": 365000, "WBS Element": "WOIN.001", "Start Date for Period of Performance": datetime(2026, 4, 1),
            "End Date for Period of Performance": datetime(2027, 3, 31), "Deletion Indicator": None,
            "PO Net Price in Group Currency": 365000}
    base.update(kw)
    return [base.get(h) for h in ZH[:-3]] + [None, base.get("GR Amount In LC"), base.get("Nature")]


def zmm_book(rows, sheet="ZMM_PO_Report", wb=None):
    own = wb is None
    wb = wb or openpyxl.Workbook()
    ws = wb.active if own else wb.create_sheet(sheet)
    ws.title = sheet
    ws.append(["Total"] + [None] * (len(ZH) - 1))
    ws.append(ZH)
    for r in rows:
        ws.append(r)
    if own:
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()
    return wb


BASE_ZMM = [
    zmm_row(4200000001, 10, **{"Net Order Value": 365000, "PO Net Price in Group Currency": 365000}),     # renewal of line 1
    zmm_row(4200000002, 10, Supplier=1000002, **{"Supplier Name": "Vendor B", "Net Order Value": 730000,
                                                 "PO Net Price in Group Currency": 730000}),         # shared by lines 2 and 3
    zmm_row(4200000009, 10, **{"WBS Element": "WOIN.777", "Supplier Name": "New Vendor", "Supplier": 1000009}),  # new PO
    zmm_row(4200000010, 10, **{"WBS Element": "WCIN.001"}),                                          # capex
    zmm_row(4200000011, 10, **{"Deletion Indicator": "L"}),                                          # deleted → not required
]


def tracker_book():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Opex_Forecast"
    h = ["S. No.", "AOP Code", "WBS Element", "Reporting Tag", "Category - 1 (CA/CR/Others)", "P&L Head", "Old PO_Unique",
         "Supplier Code", "Vendor Name", "PO Start Date", "PO End Date", "PO AMOUNT", "Net PO", "PO Nature",
         "Nature of Expense", "Budgeted FY'27", "New PO(s) mapped (ZMM)", "Head"]
    ws.append(h)
    ws.append([1, "OPA1", "WOIN.001", "DIAL", "CA", "India", 1100000001, 1000001, "Vendor A", datetime(2025, 4, 1),
               datetime(2026, 3, 31), 365000, 365000, "Recurring", "AMC & CMC", 365000, "4200000001", "x"])
    ws.append([2, "OPA2", "WOIN.002", "GHIAL", "CA", "India", 1100000002, 1000002, "Vendor B", datetime(2025, 4, 1),
               datetime(2026, 3, 31), 100000, 100000, "Recurring", "AMC & CMC", 100000, "4200000002", None])
    ws.append([3, "OPA3", "WOIN.003", "GHIAL", "CA", "India", 1100000003, 1000002, "Vendor B", datetime(2025, 4, 1),
               datetime(2026, 3, 31), 300000, 300000, "Recurring", "AMC & CMC", 300000, "4200000002", None])
    ws.append([4, "OPA4", "WOIN.004", "DIAL", "CA", "India", "AGREEMENT-ARINC 2", None, "Arinc", datetime(2025, 4, 1),
               datetime(2026, 3, 31), 50000, 50000, "Recurring", "AMC & CMC", 50000, "ARINC", None])
    ws.append([None, "OPA5", "WOIN.005", "DIAL", "CA", "India", "PROC/WAISL/21-22/363", None, "X", None, None, 1, 1,
               "One-Time", "Spares", 0, None, None])  # no S. No. → rejected
    zmm_book(BASE_ZMM, wb=wb)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def upload_zmm(client, admin, rows, name="zmm.xlsx"):
    r = client.post("/api/aop/import/zmm", headers=admin, files={"file": (name, zmm_book(rows), XLSX)})
    assert r.status_code == 200, r.text
    return r.json()


def line(client, admin, lid):
    r = client.get(f"/api/aop/datasets/opex_tracker/rows?keys={lid}", headers=admin).json()
    return r["rows"][0]["fields"]


def fy(f, fy_="27"):
    return round(sum(v for k, v in f.items() if k.startswith(f"F{fy_}__")))


def test_review_end_to_end(client, admin):
    # ---------------- one-time import (§11)
    r = client.post("/api/aop/import/opex-workbook", headers=admin, files={"file": ("opex.xlsx", tracker_book(), XLSX)})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["counts"]["opex_tracker"] == 4 and out["counts"]["po_links"] == 3
    assert any("no S. No." in x for x in out["rejected"])                        # row without S. No. rejected with its row
    assert out["zmm"]["status"] == "processed"
    assert any("ignored columns" in w and "Head" in w for w in out["warnings"])
    l1, l2, l3, l4 = (line(client, admin, f"TRK-0000{i}") for i in range(1, 5))
    assert l4["mapping_status"] == "Agreement / outside SAP"
    assert l1["latest_po"] == "4200000001" and l1["previous_po"] == "1100000001" and l1["active_po_count"] == 1
    assert fy(l1) == 365000                                                      # new PO covers the whole FY
    # 4200000002 (one item) feeds lines 2 and 3 → auto split by budget 1:3, no double counting
    assert fy(l2) == 182500 and fy(l3) == 547500
    checks = client.get("/api/aop/review/checks", headers=admin).json()["rows"]
    assert sum(1 for c in checks if c["kind"] == "Auto-split allocation") == 2
    # first run accepted everything: no PO changes raised
    assert client.get("/api/aop/review/changes", headers=admin).json()["rows"] == []

    # ---------------- To map (§9.2): pre-filled types, deleted → not required, linked POs absent
    tm = client.get("/api/aop/review/to-map", headers=admin).json()["rows"]
    pos = {x["po"]: x for x in tm}
    assert set(pos) == {"4200000009", "4200000010"}
    assert pos["4200000009"]["suggested_type"] == "Opex" and pos["4200000010"]["suggested_type"] == "Capex"
    s = client.get("/api/aop/review/summary", headers=admin).json()
    assert s["to_map"] == 2 and s["changes"] == 0

    # ---------------- idempotency: the same file twice changes nothing
    before = {lid: fy(line(client, admin, lid)) for lid in ("TRK-00001", "TRK-00002", "TRK-00003")}
    run = upload_zmm(client, admin, BASE_ZMM)
    assert run["changes_flagged"] == 0 and run["forecast_delta"] == 0
    assert {lid: fy(line(client, admin, lid)) for lid in before} == before

    # ---------------- GRN-only change → nothing flagged
    grn = [zmm_row(4200000001, 10, **{"MIGO No.": 5001, "MIGO Line Item No.": 1, "GRN Amount in Group Currency": 30000,
                                      "GRN Posting Date": datetime(2026, 5, 1)})] + BASE_ZMM[1:]
    assert upload_zmm(client, admin, grn)["changes_flagged"] == 0
    # ---------------- period change → flagged, forecast held until accepted
    moved = [zmm_row(4200000001, 10, **{"End Date for Period of Performance": datetime(2026, 9, 30)})] + BASE_ZMM[1:]
    run = upload_zmm(client, admin, moved)
    assert run["changes_flagged"] == 1
    assert fy(line(client, admin, "TRK-00001")) == 365000
    ch = client.get("/api/aop/review/changes", headers=admin).json()["rows"]
    assert len(ch) == 1 and ch[0]["field"] == "period_end" and ch[0]["po_fy_impact"] > 0
    # revert in SAP → the change closes itself
    upload_zmm(client, admin, BASE_ZMM)
    assert client.get("/api/aop/review/changes", headers=admin).json()["rows"] == []
    # flag again and accept → re-phased
    upload_zmm(client, admin, moved)
    ch = client.get("/api/aop/review/changes", headers=admin).json()["rows"]
    r = client.post("/api/aop/review/changes/decide", headers=admin, json={"keys": [ch[0]["key"]], "accept": True})
    assert r.status_code == 200 and r.json()["decided"] == 1
    l1 = line(client, admin, "TRK-00001")
    assert fy(l1) > 365000 and l1["F27__2026-10"] > 0       # same value over Apr–Sep, then the budget gap fill after Sep

    # ---------------- reject → correction with the accepted value as override; fixed in SAP → Possibly resolved
    val = [zmm_row(4200000001, 10, **{"End Date for Period of Performance": datetime(2026, 9, 30),
                                      "Net Order Value": 999000, "PO Net Price in Group Currency": 999000})] + BASE_ZMM[1:]
    upload_zmm(client, admin, val)
    ch = client.get("/api/aop/review/changes", headers=admin).json()["rows"]
    assert {c["field"] for c in ch} == {"net_order_value"}
    assert client.post("/api/aop/review/changes/decide", headers=admin, json={"po": "4200000001", "accept": False}).status_code == 400
    r = client.post("/api/aop/review/changes/decide", headers=admin,
                    json={"po": "4200000001", "accept": False, "remarks": "Wrong value in SAP"})
    assert r.json()["corrections"] == 1
    cors = client.get("/api/aop/review/corrections", headers=admin).json()["rows"]
    cor = next(c for c in cors if c["source"] == "rejected change")
    assert cor["type"] == "Wrong value / qty" and cor["override"]["net_order_value"] == 365000
    # same wrong value next run → not raised again
    assert upload_zmm(client, admin, val)["changes_flagged"] == 0
    fixed = [zmm_row(4200000001, 10, **{"End Date for Period of Performance": datetime(2026, 9, 30),
                                        "Net Order Value": 400000, "PO Net Price in Group Currency": 400000})] + BASE_ZMM[1:]
    upload_zmm(client, admin, fixed)
    cors = client.get("/api/aop/review/corrections", headers=admin).json()["rows"]
    assert next(c for c in cors if c["key"] == cor["key"])["status"] == "Possibly resolved"
    r = client.patch(f"/api/aop/review/corrections/{cor['key']}", headers=admin, json={"status": "Resolved"})
    assert r.status_code == 200 and r.json()["override"] is None
    x = client.get("/api/aop/review/corrections?format=xlsx", headers=admin)
    assert x.status_code == 200 and x.content[:2] == b"PK"

    # ---------------- new item on a mapped PO → change with pre-filled mapping
    added = fixed + [zmm_row(4200000001, 20, Material=9700000099, **{"Net Order Value": 10000})]
    upload_zmm(client, admin, added)
    ch = client.get("/api/aop/review/changes", headers=admin).json()["rows"]
    ia = [c for c in ch if c["field"] == "item_added"]
    assert len(ia) == 1 and ia[0]["prefill"]["line_id"] == "TRK-00001" and ia[0]["prefill"]["type"] == "Opex"
    r = client.post("/api/aop/review/changes/decide", headers=admin, json={"keys": [c["key"] for c in ch], "accept": True})
    assert r.status_code == 200

    # ---------------- To map decisions: Opex renewal as an add-on line; capex; bulk
    r = client.post("/api/aop/review/to-map/decide", headers=admin, json={"decisions": [
        {"po": "4200000009", "type": "Opex", "opex_action": "Renewal / replacement", "line_id": "TRK-00001", "add_as_new_line": True},
        {"po": "4200000010", "type": "Capex", "capex_key": "DIAL|Network"}]})
    assert r.status_code == 200 and r.json()["decided"] == 2, r.text
    addon = line(client, admin, "TRK-00001-A1")
    assert addon["parent_line_id"] == "TRK-00001" and addon["mapping_status"] == "Add-on" and addon["aop_code"] == "OPA1"
    assert addon["B27__annual"] == 0 and fy(addon) == 365000 and addon["latest_po"] == "4200000009"
    assert client.get("/api/aop/review/to-map", headers=admin).json()["rows"] == []

    # ---------------- drawers
    d = client.get("/api/aop/po/4200000002", headers=admin).json()
    assert {ln["line_id"] for ln in d["links"]} == {"TRK-00002", "TRK-00003"} and d["summary"]["po_value"] == 730000
    h = client.get("/api/aop/opex/lines/TRK-00001/history", headers=admin).json()
    assert h["timeline"][0]["own"] and h["timeline"][0]["po"] == "1100000001" and h["detail"]["kind"] == "sap"
    h = client.get("/api/aop/opex/lines/TRK-00001/history?po=1100000001", headers=admin).json()
    assert h["detail"]["kind"] == "legacy" and h["detail"]["fields"]["net_po"] == 365000
    assert client.post("/api/aop/opex/lines/TRK-00002/add-on", headers=admin, json={}).json()["line_id"] == "TRK-00002-A1"

    # ---------------- upload log
    runs = client.get("/api/aop/review/runs", headers=admin).json()["rows"]
    assert len(runs) >= 8 and all(x["status"] == "processed" for x in runs)
    f = client.get(f"/api/aop/review/runs/{runs[0]['run_id']}/file", headers=admin)
    assert f.status_code == 200 and f.content[:2] == b"PK"

    # ---------------- download → upload round trip changes nothing
    dl = client.get("/api/aop/datasets/opex_tracker/download", headers=admin)
    assert dl.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(dl.content), read_only=True)
    hdr = [c for c in next(wb.worksheets[0].iter_rows(values_only=True))]
    assert hdr[0] == "Line ID" and "Budgeted FY'28" in hdr and "Mapping status" in hdr
    before = {lid: line(client, admin, lid) for lid in ("TRK-00001", "TRK-00002")}
    up = client.post("/api/aop/datasets/opex_tracker/upload?mode=modify", headers=admin,
                     files={"file": ("opex_tracker.xlsx", dl.content, XLSX)})
    assert up.status_code == 200, up.text
    assert up.json()["ignored_columns"] == [] and up.json()["errors"] == []
    after = {lid: line(client, admin, lid) for lid in before}
    for lid in before:
        for k in ("aop_code", "po", "recurring", "B27__annual", "mapping_status"):
            assert after[lid].get(k) == before[lid].get(k), (lid, k)
        assert fy(after[lid]) == fy(before[lid])

    # ---------------- permissions: Review needs edit on Opex (or the Review section)
    s = client.get("/api/aop/review/summary", headers=admin).json()
    assert s["to_map"] == 0 and "badge" in s


def test_review_requires_opex_edit(client, make_user):
    viewer = make_user({"aop_opex": {"can_view": True}})
    assert client.get("/api/aop/review/summary", headers=viewer).status_code == 403
    editor = make_user({"aop_opex": {"can_view": True, "can_edit": True}})
    assert client.get("/api/aop/review/summary", headers=editor).status_code == 200
    assert client.post("/api/aop/import/zmm", headers=viewer,
                       files={"file": ("z.xlsx", zmm_book(BASE_ZMM), XLSX)}).status_code == 403

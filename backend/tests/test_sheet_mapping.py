"""One PO mapping edited from three places (Opex sheet, ZMM sheet, PO links) and corrections typed in the ZMM sheet.
Runs on the state test_review leaves (tracker lines TRK-00001…4, the base ZMM)."""
import io

import openpyxl

from test_review import BASE_ZMM, XLSX, ZH, zmm_row


def book(rows, extra_header=(), extra=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(list(ZH) + list(extra_header))
    for i, r in enumerate(rows):
        ws.append(list(r) + (list(extra[i]) if i < len(extra) else []))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


NEW = [zmm_row(4300000001, 10), zmm_row(4300000001, 20, Material=9700000002),
       zmm_row(4300000002, 10, **{"Supplier Name": "Vendor B", "Supplier": 1000002})]


def rows(client, admin, ds, keys=None):
    params = {"keys": chr(1).join(keys)} if keys else {"limit": 5000}
    r = client.get(f"/api/aop/datasets/{ds}/rows", headers=admin, params=params)
    assert r.status_code == 200, r.text
    return {d["key"]: d["fields"] for d in r.json()["rows"]}


def patch(client, admin, ds, edits):
    r = client.patch(f"/api/aop/datasets/{ds}/rows", headers=admin, json=edits)
    assert r.status_code == 200, r.text
    return r.json()


def links(client, admin):
    return rows(client, admin, "po_links")


def test_mapping_from_three_places_and_sheet_corrections(client, admin):
    r = client.post("/api/aop/import/zmm", headers=admin, files={"file": ("zmm.xlsx", book(BASE_ZMM + NEW), XLSX)})
    assert r.status_code == 200, r.text

    # ---------------- 1. Opex sheet: "New PO(s) mapped"
    before = rows(client, admin, "opex_tracker", ["TRK-00002"])["TRK-00002"]["mapped_pos"]
    assert "4200000002" in before
    out = patch(client, admin, "opex_tracker", [{"key": "TRK-00002", "field": "mapped_pos", "value": f"{before}, 4300000002"}])
    assert out["applied"] == 1 and not out["rejected"]
    assert "TRK-00002|4300000002||" in links(client, admin)
    assert rows(client, admin, "po_items", ["4300000002|10"])["4300000002|10"]["linked_lines"] == "TRK-00002"
    # words set the status
    patch(client, admin, "opex_tracker", [{"key": "TRK-00004", "field": "mapped_pos", "value": "Not required"}])
    assert rows(client, admin, "opex_tracker", ["TRK-00004"])["TRK-00004"]["mapping_status"] == "Not required"

    # an Opex line with no tracker line gets one; item-level token PO/item
    r = client.post("/api/aop/datasets/opex_lines/rows", headers=admin, json=[{"aop_code": "OPXA9", "tag": "DIAL", "wbs": "WOIN.009"}])
    assert r.status_code == 200, r.text
    ol = next(k for k, f in rows(client, admin, "opex_lines").items() if f.get("aop_code") == "OPXA9")
    patch(client, admin, "opex_lines", [{"key": ol, "field": "mapped_pos", "value": "4300000001/20"}])
    f = rows(client, admin, "opex_lines", [ol])[ol]
    tid = f["tracker_line_id"]
    assert tid.startswith("TRK-") and f["mapped_pos"] == "4300000001/20"
    assert rows(client, admin, "opex_tracker", [tid])[tid]["opex_line_id"] == ol
    its = rows(client, admin, "po_items", ["4300000001|10", "4300000001|20"])
    assert its["4300000001|20"]["linked_lines"] == tid and not its["4300000001|10"]["linked_lines"]

    # ---------------- 2. ZMM sheet: "Opex line(s)" by S. No. / Opex line id; clearing removes
    patch(client, admin, "po_items", [{"key": "4300000001|10", "field": "linked_lines", "value": "1"}])
    assert "TRK-00001|4300000001||10" in links(client, admin)  # two items → item-level link
    patch(client, admin, "po_items", [{"key": "4300000001|20", "field": "linked_lines", "value": f"{ol}, 3"}])
    assert rows(client, admin, "po_items", ["4300000001|20"])["4300000001|20"]["linked_lines"] == f"{tid}, TRK-00003"
    patch(client, admin, "po_items", [{"key": "4300000002|10", "field": "linked_lines", "value": ""}])
    assert "TRK-00002|4300000002||" not in links(client, admin)
    assert "4300000002" not in (rows(client, admin, "opex_tracker", ["TRK-00002"])["TRK-00002"]["mapped_pos"] or "")
    bad = patch(client, admin, "po_items", [{"key": "4300000002|10", "field": "linked_lines", "value": "NOPE-1"}])
    assert bad["applied"] == 0 and "no Opex line" in bad["rejected"][0]["reason"]

    # ---------------- 3. corrections typed over SAP's value
    out = patch(client, admin, "po_items", [{"key": "4300000001|10", "field": "period_end", "value": "2027-06-30"}])
    assert out["correction"] == 1
    it = rows(client, admin, "po_items", ["4300000001|10"])["4300000001|10"]
    assert it["period_end"] == "2027-06-30" and "SAP 2027-03-31 → 2027-06-30" in it["corrected"]
    cors = {k: f for k, f in rows(client, admin, "po_corrections").items() if f.get("po") == "4300000001"}
    assert len(cors) == 1 and next(iter(cors.values()))["override"] == {"period_end": "2027-06-30"}
    # same correction gets the WBS too; typing SAP's value back drops the field, the last one resolves it
    patch(client, admin, "po_items", [{"key": "4300000001|10", "field": "wbs", "value": "WOIN.900"}])
    cid, c = next((k, f) for k, f in rows(client, admin, "po_corrections").items() if f.get("po") == "4300000001")
    assert c["override"] == {"period_end": "2027-06-30", "wbs": "WOIN.900"} and c["type"] == "Other"
    patch(client, admin, "po_items", [{"key": "4300000001|10", "field": "period_end", "value": "2027-03-31"},
                                      {"key": "4300000001|10", "field": "wbs", "value": ""}])
    c = rows(client, admin, "po_corrections", [cid])[cid]
    assert c["status"] == "Resolved" and not c.get("override")
    # the raw ZMM sheet (po_register) takes the same edits under its SAP headers
    reg = next(k for k, f in rows(client, admin, "po_register").items() if f.get("purchase_order") == "4300000002")
    out = patch(client, admin, "po_register", [{"key": reg, "field": "wbs_element", "value": "WOIN.555"},
                                               {"key": reg, "field": "linked_forecast_s_no", "value": "2"}])
    assert out == {**out, "applied": 2, "correction": 1, "mapping": 1}
    f = rows(client, admin, "po_register", [reg])[reg]
    assert f["wbs_element"] == "WOIN.555" and f["linked_forecast_s_no"] == "TRK-00002" and "SAP WOIN.001" in f["corrected"]
    assert rows(client, admin, "po_register", [reg])
    nope = patch(client, admin, "po_register", [{"key": reg, "field": "supplier_name_2", "value": "x"}])
    assert nope["applied"] == 0

    # ---------------- 4. the ZMM sheet downloaded, edited in Excel, uploaded back
    x = client.get("/api/aop/datasets/po_items/download?fmt=xlsx", headers=admin)
    assert x.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    ws = wb.active
    head = [c.value for c in ws[1]]
    ci, cp, ce = head.index("linked_lines"), head.index("po"), head.index("period_start")
    for row in ws.iter_rows(min_row=2):
        if str(row[cp].value) == "4300000002":
            assert row[ci].value == "TRK-00002"
            row[ci].value = "TRK-00002, 4"
            row[ce].value = "2026-05-01"
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/aop/datasets/po_items/upload", headers=admin, files={"file": ("po_items.xlsx", buf.getvalue(), XLSX)})
    assert r.status_code == 200, r.text
    assert r.json()["mapping"] == 1 and r.json()["correction"] == 1 and r.json()["applied"] == 2, str(r.json())
    it = rows(client, admin, "po_items", ["4300000002|10"])["4300000002|10"]
    assert it["linked_lines"] == "TRK-00002, TRK-00004" and it["period_start"] == "2026-05-01"

    # ---------------- 5. a ZMM file that carries the mapping ("Linked Forecast S.No") adds it
    more = NEW + [zmm_row(4300000003, 10)]
    r = client.post("/api/aop/import/zmm", headers=admin, files={"file": ("zmm2.xlsx", book(
        BASE_ZMM + more, ["Linked Forecast S.No"], [()] * len(BASE_ZMM) + [(), (), (), (3,)]), XLSX)})
    assert r.status_code == 200, r.text
    assert r.json()["links_from_file"] == 1, r.json()
    assert "TRK-00003|4300000003||" in links(client, admin)
    # corrections survive the next SAP report (SAP still has the old value)
    it = rows(client, admin, "po_items", ["4300000002|10"])["4300000002|10"]
    assert it["period_start"] == "2026-05-01" and "SAP 2026-04-01" in it["corrected"]

    # ---------------- 6. the PO links tab still edits directly, and the other two places show it
    r = client.request("DELETE", "/api/aop/datasets/po_links/rows", headers=admin, json=["TRK-00003|4300000003||"])
    assert r.status_code == 200, r.text
    assert not rows(client, admin, "po_items", ["4300000003|10"])["4300000003|10"]["linked_lines"]


def test_forecast_file_gives_mapping_and_corrections_without_rolling_back_the_zmm(client, admin):
    from datetime import datetime
    from test_review import tracker_book
    wb = openpyxl.load_workbook(io.BytesIO(tracker_book()))
    ws = wb["ZMM_PO_Report"]
    ws.delete_rows(3, ws.max_row)
    ws.cell(2, len(ZH) + 1, "Linked Forecast S.No")
    old = {"Created On": datetime(2026, 1, 1)}
    for r in [zmm_row(4200000001, 10, **old, **{"End Date for Period of Performance": datetime(2027, 9, 30)}),  # corrected
              zmm_row(4300000003, 10, **old)]:
        ws.append(r + ([1] if r[0] == 4300000003 else [None]))
    buf = io.BytesIO()
    wb.save(buf)
    before = rows(client, admin, "po_items")
    r = client.post("/api/aop/import/opex-workbook", headers=admin, files={"file": ("opex.xlsx", buf.getvalue(), XLSX)})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["zmm"] is None and out["counts"]["zmm_links"] == 1 and out["counts"]["zmm_corrections"] == 1
    assert any("older than the loaded ZMM" in w for w in out["warnings"])
    assert set(rows(client, admin, "po_items")) == set(before)  # the loaded SAP report is kept
    it = rows(client, admin, "po_items", ["4200000001|10"])["4200000001|10"]
    assert it["period_end"] == "2027-09-30" and "SAP 2027-03-31" in it["corrected"]
    assert "TRK-00001|4300000003||" in links(client, admin)
    cor = next(f for f in rows(client, admin, "po_corrections").values() if f.get("po") == "4200000001" and f.get("status") == "Open")
    assert cor["source"] == "Opex forecast file" and cor["override"] == {"period_end": "2027-09-30"}

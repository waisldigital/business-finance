"""ZMM pipeline, link resolution and the tracker recalculation (database side of ``po.py``).

One pipeline for both entry points — the manual upload (``POST /api/aop/import/zmm``) and the scheduled e-mail fetch
(``jobs/fetch_zmm.py``)::

    run_zmm_pipeline(content, filename, source, meta)
      1 parse, validate, normalise; store the file; open a zmm_runs row
      2 replace po_register
      3 upsert po_triage enrichment; auto-triage (deleted / blocked → Mapping not required; type from the WBS prefix)
      4 rebuild po_items latest_* values (accepted_* kept for items in review)
      5 detect changes on mapped items → po_changes
      6 resolve all links on accepted values      7 recalc the forecast and tracker display fields
      8 write the derived ZMM columns back onto po_register      9 close the run, notify

The same file processed twice changes nothing.
"""
from __future__ import annotations

import bisect
import hashlib
import io
import logging
import os
from collections import defaultdict
from datetime import datetime, timezone
from typing import Any, Awaitable, Callable, Dict, List, Optional, Tuple

import openpyxl
from pymongo import UpdateOne

from . import po as P
from .datasets import column, vkey
from .opex import forecast_line
from .opex_schema import ZMM_KEYS
from .periods import fy_months

logger = logging.getLogger("finsight.aop.po")

DEFAULT_FX = {"EUR": 108.0, "USD": 90.0, "GBP": 110.0}  # Mapping!K7:L11 of the forecast workbook
RUN_RETENTION_YEARS = 3


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _num(v) -> float:
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


FX_API = os.environ.get("FX_API_URL", "https://api.frankfurter.dev/v1")
# All-currency daily rates (incl. AED, SAR, QAR … that the ECB doesn't publish), one file per date; two mirrors.
FX_API_ALT = [u for u in os.environ.get(
    "FX_API_ALT_URLS", "https://cdn.jsdelivr.net/npm/@fawazahmed0/currency-api@{date}/v1/currencies/{cur}.json,"
                       "https://{date}.currency-api.pages.dev/v1/currencies/{cur}.json").split(",") if u.strip()]
FX_COMMON = ["USD", "EUR", "GBP", "AED", "SGD"]
FX_MAX_ALT_DAYS = 400


def fx_auto_fetch() -> bool:
    return os.environ.get("FX_AUTO_FETCH", "true").strip().lower() not in ("0", "false", "no", "off")


async def _fetch_ecb(c, cur: str, lo: str, end: str) -> Dict[str, float]:
    r = await c.get(f"{FX_API}/{lo}..{end}", params={"from": cur, "to": "INR"})
    if r.status_code in (404, 422):  # currency not published by the ECB
        return {}
    r.raise_for_status()
    return {dt: float(v["INR"]) for dt, v in (r.json().get("rates") or {}).items() if v.get("INR")}


async def _fetch_alt(c, cur: str, day: str) -> Optional[float]:
    for tpl in FX_API_ALT:
        try:
            r = await c.get(tpl.strip().format(date=day, cur=cur.lower()))
            if r.status_code == 200:
                v = (r.json().get(cur.lower()) or {}).get("inr")
                if v:
                    return float(v)
        except Exception:  # noqa: BLE001 — try the next mirror
            continue
    return None


async def fetch_fx_series(cur: str, start: str, end: str, dates: Optional[List[str]] = None) -> Tuple[Dict[str, float], str]:
    """INR per unit of ``cur`` by date for [start − 10 days, end]: ECB reference rates (business days); for a
    currency the ECB doesn't publish, the daily all-currency rates for ``dates`` (else every day in the range)."""
    import httpx
    from datetime import date as _d, timedelta
    lo = (_d.fromisoformat(start) - timedelta(days=10)).isoformat()
    async with httpx.AsyncClient(timeout=30, follow_redirects=True) as c:
        out = await _fetch_ecb(c, cur, lo, end)
        if out:
            return out, "ECB reference (auto)"
        days = sorted(set(dates or [])) or [(_d.fromisoformat(start) + timedelta(days=i)).isoformat()
                                            for i in range((_d.fromisoformat(end) - _d.fromisoformat(start)).days + 1)]
        today = _d.today().isoformat()
        for day in [x for x in days if x <= today][-FX_MAX_ALT_DAYS:]:
            v = await _fetch_alt(c, cur, day)
            if v:
                out[day] = v
    return out, "Daily market rate (auto)"


class FxTable:
    """Rate on a date = that day's rate, else the nearest earlier one within 10 days (weekends, holidays)."""

    def __init__(self, stored: Dict[str, Dict[str, float]], fallback: Dict[str, Optional[float]]):
        self.stored, self.fallback = stored, fallback
        self.fetched, self.error = 0, None
        self.reindex()

    def reindex(self):
        self.idx = {c: sorted(v.items()) for c, v in self.stored.items()}

    def lookup(self, cur: str, dt: str) -> Optional[Tuple[float, str]]:
        series = self.idx.get(cur) or sorted((self.stored.get(cur) or {}).items())
        if not series or not dt:
            return None
        i = bisect.bisect_right([x[0] for x in series], dt) - 1
        if i < 0:
            return None
        day, rate = series[i]
        from datetime import date as _d
        if (_d.fromisoformat(dt) - _d.fromisoformat(day)).days > 10:
            return None
        return rate, f"FX {cur} on PO date {day}"

    def rate(self, cur: str, po_date) -> Optional[Tuple[float, str]]:
        got = self.lookup(cur, po_date.isoformat() if po_date else "")
        if got:
            return got
        fb = self.fallback.get(cur)
        return (fb, P.FX_FALLBACK) if fb else None


class OpexEngine:
    """Async helpers over ``aop_rows`` for the PO datasets. ``get_config`` returns the aop_config document;
    ``fx_rate(currency)`` → INR per unit (Assumptions); ``storage`` (optional) keeps the uploaded files."""

    def __init__(self, db, get_config: Callable[[], Awaitable[Dict[str, Any]]],
                 fx_rate: Callable[[str], Awaitable[Optional[float]]], storage=None,
                 bump: Optional[Callable[[], Awaitable[None]]] = None):
        self.db = db
        self.get_config = get_config
        self.fx_rate = fx_rate
        self.storage = storage
        self.bump = bump

    # ------------------------------------------------------------------ small helpers
    async def rows(self, dataset: str, q: Optional[Dict[str, Any]] = None) -> List[Dict[str, Any]]:
        mq = {"dataset": dataset, **(q or {})}
        return [d async for d in self.db.aop_rows.find(mq, {"_id": 0, "key": 1, "fields": 1, "seq": 1})]

    async def fields(self, dataset: str, q: Optional[Dict[str, Any]] = None) -> Dict[str, Dict[str, Any]]:
        return {d["key"]: d.get("fields") or {} for d in await self.rows(dataset, q)}

    async def bulk_set(self, dataset: str, updates: Dict[str, Dict[str, Any]], *, upsert: bool = False, by: str = "system",
                       unset: Optional[Dict[str, List[str]]] = None):
        ops = []
        seq = None
        for key, upd in updates.items():
            s: Dict[str, Any] = {f"fields.{k}": v for k, v in upd.items()}
            s["updated_at"] = now_iso()
            s["updated_by"] = by
            op: Dict[str, Any] = {"$set": s}
            if unset and unset.get(key):
                op["$unset"] = {f"fields.{k}": "" for k in unset[key]}
            if upsert:
                if seq is None:
                    last = await self.db.aop_rows.find_one({"dataset": dataset}, {"_id": 0, "seq": 1}, sort=[("seq", -1)])
                    seq = (last or {}).get("seq") or 0
                seq += 1
                op["$setOnInsert"] = {"seq": seq}
            ops.append(UpdateOne({"dataset": dataset, "key": key}, op, upsert=upsert))
        for i in range(0, len(ops), 1000):
            await self.db.aop_rows.bulk_write(ops[i:i + 1000], ordered=False)

    async def rate_map(self) -> Dict[str, Optional[float]]:
        out = {}
        for cur in ("USD", "EUR", "GBP", "AED", "SGD", "AUD", "CAD", "CHF", "JPY"):
            out[cur] = await self.fx_rate(cur) or DEFAULT_FX.get(cur)
        return out

    async def fx_table(self, rows: List[Dict[str, Any]], by: str = "system") -> "FxTable":
        """Exchange rates for the PO dates in this file: the stored FX table (aop_rows fx_rates, editable / uploadable
        under Inputs), missing dates fetched from the ECB reference rates (FX_AUTO_FETCH, default on) and saved, and
        the Assumptions FY rate as the last resort (flagged)."""
        stored: Dict[str, Dict[str, float]] = defaultdict(dict)
        for f in (await self.fields("fx_rates")).values():
            cur, dt, rate = str(f.get("currency") or "").upper(), str(f.get("date") or "")[:10], f.get("rate")
            if cur and dt and isinstance(rate, (int, float)) and rate > 0:
                stored[cur][dt] = float(rate)
        need: Dict[str, List[str]] = defaultdict(list)
        for f in rows:
            cur = str(f.get("currency") or "INR").strip().upper()
            dt = str(f.get("created_on") or "")[:10]
            if cur not in ("", "INR") and dt:
                need[cur].append(dt)
        table = FxTable(stored, await self.rate_map())
        missing = {c: [x for x in ds if table.lookup(c, x) is None] for c, ds in need.items()}
        missing = {c: ds for c, ds in missing.items() if ds}
        if missing and fx_auto_fetch():
            res = await self.fetch_fx(missing, by=by, stored=stored)
            table.fetched, table.error = res["saved"], "; ".join(res["errors"]) or None
        table.reindex()
        return table

    async def fetch_fx(self, need: Dict[str, List[str]], by: str = "system",
                       stored: Optional[Dict[str, Dict[str, float]]] = None, overwrite: bool = False) -> Dict[str, Any]:
        """Fetch rates from the internet for ``{currency: [dates]}`` (the range min..max of the dates is fetched) and
        save them in fx_rates. Rates typed or uploaded by hand (source not "… (auto)") are never overwritten."""
        cur_rows = await self.fields("fx_rates")
        manual = {k for k, f in cur_rows.items() if f.get("source") and "(auto)" not in str(f.get("source"))}
        new: Dict[str, Dict[str, Any]] = {}
        errors: List[str] = []
        per_cur: Dict[str, int] = {}
        for cur, ds in need.items():
            cur = str(cur).upper()
            if cur in ("", "INR") or not ds:
                continue
            try:
                got, src = await fetch_fx_series(cur, min(ds), max(ds), ds)
            except Exception as e:  # noqa: BLE001 — no network / API down: callers fall back to the assumptions
                errors.append(f"{cur}: {e}"[:300])
                continue
            if not got:
                errors.append(f"{cur}: no rates published")
            n = 0
            for dt, rate in got.items():
                k = f"{cur}|{dt}"
                if k in manual or (k in cur_rows and (not overwrite or cur_rows[k].get("rate") == rate)):
                    continue
                if stored is not None:
                    stored.setdefault(cur, {})[dt] = rate
                new[k] = {"currency": cur, "date": dt, "rate": rate, "source": src}
                n += 1
            per_cur[cur] = n
        if new:
            await self.bulk_set("fx_rates", new, upsert=True, by=by)
            if self.bump:
                await self.bump()
        return {"saved": len(new), "by_currency": per_cur, "errors": errors}

    async def po_fx_need(self, start: Optional[str] = None, end: Optional[str] = None,
                         currencies: Optional[List[str]] = None) -> Dict[str, List[str]]:
        """What to fetch: every foreign currency on the PO items (else the common ones) for their PO dates;
        with ``start``/``end`` the whole range instead (default: the last 30 days up to today)."""
        from datetime import date as _d, timedelta
        po_dates: Dict[str, List[str]] = defaultdict(list)
        for f in (await self.fields("po_items")).values():
            cur = str(f.get("currency") or "INR").strip().upper()
            dt = str(f.get("created_on") or "")[:10]
            if cur not in ("", "INR") and dt:
                po_dates[cur].append(dt)
        curs = [c.upper() for c in (currencies or [])] or sorted(po_dates) or FX_COMMON
        if start or end or not po_dates:
            hi = end or _d.today().isoformat()
            lo = start or (_d.fromisoformat(hi) - timedelta(days=30)).isoformat()
            n = (_d.fromisoformat(hi) - _d.fromisoformat(lo)).days
            if n < 0:
                raise ValueError("From date is after To date")
            days = [(_d.fromisoformat(lo) + timedelta(days=i)).isoformat() for i in range(min(n, FX_MAX_ALT_DAYS) + 1)]
            return {c: days for c in curs}
        today = _d.today().isoformat()
        return {c: sorted(set(po_dates.get(c) or [])) + [today] for c in curs}

    async def corrections_overrides(self) -> Dict[str, Dict[str, Any]]:
        out: Dict[str, Dict[str, Any]] = {}
        for f in (await self.fields("po_corrections")).values():
            if f.get("status") in ("Resolved",) or not f.get("override"):
                continue
            for it in f.get("items") or []:
                out.setdefault(f"{f.get('po')}|{it}", {}).update(f["override"])
        return out

    # ------------------------------------------------------------------ resolution + forecast
    async def load_state(self) -> Dict[str, Any]:
        cfg = await self.get_config()
        raw_items = await self.fields("po_items")
        overrides = await self.corrections_overrides()
        return {"cfg": cfg, "raw_items": raw_items, "overrides": overrides,
                "items": {k: P.effective(v, overrides.get(k)) for k, v in raw_items.items()},
                "links": [{**f, "_key": k} for k, f in (await self.fields("po_links")).items()],
                "lines": await self.fields("opex_tracker"),
                "open_changes": list((await self.fields("po_changes", {"fields.status": "Open"})).values())}

    def compute(self, st: Dict[str, Any], *, items: Optional[Dict[str, Dict[str, Any]]] = None,
                only: Optional[set] = None, details: bool = True) -> Dict[str, Any]:
        """Links → cover, allocation, segments and the forecast for every line (in memory). ``items`` replaces the
        effective item values (dry run with SAP's latest values); ``only`` limits the per-line output."""
        cfg = st["cfg"]
        plan = cfg["plan_fy"]
        items = items if items is not None else st["items"]
        links, lines = st["links"], st["lines"]
        cover = P.resolve(links, items)
        weights = {lid: {"budget": _num(f.get(vkey("B" + plan[2:], "annual"), f.get("budget_plan"))),
                         "net_po": _num(f.get("net_po") or f.get("po_amount"))} for lid, f in lines.items()}
        alloc = P.allocate(links, cover, weights, self._history(links, cover, items))
        by_line: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
        for ln in links:
            by_line[ln["line_id"]].append(ln)
        open_changes = defaultdict(int)
        for c in st["open_changes"]:
            open_changes[c.get("po")] += 1
        own_items: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
        for it in items.values():
            own_items[it["po"]].append(it)
        line_upd: Dict[str, Dict[str, Any]] = {}
        forecasts: Dict[str, float] = {}
        for lid, line in lines.items():
            if only is not None and lid not in only:
                continue
            lks = by_line.get(lid, [])
            fill = self._zmm_fill(line, own_items.get(str(line.get("po") or ""), []))
            row = {**line, **fill}
            segs, detail, flags = P.segments(row, lks, cover, items, alloc)
            fc = forecast_line(row, plan, segs if lks else None)
            forecasts[lid] = sum(fc.values())
            if not details:
                continue
            lp = P.latest_previous(row, detail)
            lp.pop("po_order", None)
            for ln in lks:
                flags += [f"{ln['po']}: {x}" for x in cover.get(ln["_key"], {}).get("flags", [])]
            mm = P.rate_mismatch(row, detail)
            upd = {**fill, **fc, **lp,
                   "new_po_fy_inr": round(P.fy_amount(segs, plan), 2),
                   "new_po_amount": round(sum(x["amount"] for x in detail), 2) if detail else None,
                   "link_flags": "; ".join(dict.fromkeys(flags + ([mm] if mm else []))) or None,
                   "pending_review": sum(open_changes.get(p, 0) for p in {ln["po"] for ln in lks}) or None}
            if not lks:
                upd.update(latest_po=line.get("po"), previous_po=None, active_po_count=0)
            line_upd[lid] = upd
        return {"cover": cover, "alloc": alloc, "line_upd": line_upd, "forecasts": forecasts, "items": items}

    async def resolve_all(self, line_keys: Optional[List[str]] = None) -> Dict[str, Any]:
        """Resolve every link on effective (accepted) item values, apply auto allocation, and write the tracker
        lines' forecast and PO display fields (all lines when line_keys is None) plus the links' computed fields."""
        st = await self.load_state()
        out = self.compute(st, only=set(line_keys) if line_keys is not None else None)
        cover, alloc, items, lines = out["cover"], out["alloc"], out["items"], st["lines"]
        link_upd: Dict[str, Dict[str, Any]] = {}
        for ln in st["links"]:
            c = cover.get(ln["_key"], {})
            a = alloc.get(ln["_key"], {})
            cov_items = [items[k] for k in c.get("items", [])]
            pct = a.get("alloc_pct", ln.get("alloc_pct"))
            share = 1.0 if pct in (None, "") else float(pct)
            val = sum(_num(i.get("value_inr")) for i in cov_items)
            starts = [i.get("period_start") for i in cov_items if i.get("period_start")]
            ends = [i.get("period_end") for i in cov_items if i.get("period_end")]
            first = cov_items[0] if cov_items else {}
            line = lines.get(ln["line_id"]) or {}
            sup_changed = None
            if first.get("supplier_code") and line.get("supplier_code"):
                sup_changed = "Yes" if first["supplier_code"] != str(line["supplier_code"]) else "No"
            upd = {"items": [k.split("|", 1)[1] for k in c.get("items", [])], "value_inr": round(val, 2),
                   "alloc_value_inr": round(val * share, 2),
                   "period_start": min(starts) if starts else None, "period_end": max(ends) if ends else None,
                   "supplier_code": first.get("supplier_code"), "supplier_name": first.get("supplier_name"),
                   "supplier_changed": sup_changed,
                   "grn_inr": round(sum(_num(i.get("grn_inr")) for i in cov_items), 2),
                   "pending_inr": round(sum(_num(i.get("pending_inr")) for i in cov_items), 2),
                   "fy_impact_inr": round(sum(_num(i.get("fy_impact_inr")) for i in cov_items) * share, 2),
                   "flags": "; ".join(c.get("flags", []) + a.get("checks", [])) or None, "resolved_at": now_iso()}
            if "alloc_pct" in a:
                upd.update(alloc_pct=a["alloc_pct"], alloc_auto=True)
            link_upd[ln["_key"]] = upd
        await self.bulk_set("opex_tracker", out["line_upd"])
        await self.bulk_set("po_links", link_upd)
        if self.bump:
            await self.bump()
        return {"lines": len(out["line_upd"]), "links": len(link_upd),
                "forecast_total": round(sum(out["forecasts"].values()), 2), "forecasts": out["forecasts"]}

    async def store_change_impacts(self):
        """FY impact of every open PO change (dry run per PO) — shown in Review and the tracker banner."""
        st = await self.load_state()
        if not st["open_changes"]:
            return
        base = self.compute(st, details=False)["forecasts"]
        by_po: Dict[str, set] = defaultdict(set)
        for c in st["open_changes"]:
            by_po[c.get("po")].add(f"{c.get('po')}|{c.get('item')}")
        impact = {po: self.change_impact(st, sorted(keys), base)[0] for po, keys in by_po.items()}
        upd = {k: {"fy_impact": impact.get(c.get("po"))}
               for k, c in (await self.fields("po_changes", {"fields.status": "Open"})).items()}
        await self.bulk_set("po_changes", upd)

    def change_impact(self, st: Dict[str, Any], keys: List[str], base: Optional[Dict[str, float]] = None) -> Tuple[float, Dict[str, float]]:
        """FY forecast impact of taking SAP's latest values for these items (dry run): (Σ Δ, {line: Δ})."""
        if base is None:
            base = self.compute(st, details=False)["forecasts"]
        items = dict(st["items"])
        for k in keys:
            raw = st["raw_items"].get(k)
            if not raw:
                continue
            latest = dict(raw)
            for f in P.TRACKED:
                if f"latest_{f}" in raw:
                    latest[f"accepted_{f}"] = raw[f"latest_{f}"]
            items[k] = P.effective(latest, st["overrides"].get(k))
        pos = {k.split("|", 1)[0] for k in keys}
        affected = {ln["line_id"] for ln in st["links"] if str(ln.get("po")) in pos} | \
            {lid for lid, f in st["lines"].items() if str(f.get("po") or "") in pos}
        after = self.compute(st, items=items, only=affected, details=False)["forecasts"]
        delta = {lid: round(after.get(lid, 0) - base.get(lid, 0), 2) for lid in after}
        delta = {k: v for k, v in delta.items() if abs(v) >= 1}
        return round(sum(delta.values()), 2), delta

    @staticmethod
    def _zmm_fill(line: Dict[str, Any], own: List[Dict[str, Any]]) -> Dict[str, Any]:
        """The line's own PO is an SAP PO in the ZMM → its Input / ZMM fields come from the PO items (by material
        when the line names one), the uploaded values only for legacy POs."""
        if not own:
            return {}
        its = [i for i in own if i.get("active", True)] or own
        if line.get("material"):
            its = [i for i in its if i.get("material") == str(line["material"])] or its
        h = its[0]
        starts = [i["period_start"] for i in its if i.get("period_start")]
        ends = [i["period_end"] for i in its if i.get("period_end")]
        out = {"po_date": h.get("created_on"), "supplier_code": h.get("supplier_code"), "supplier_name": h.get("supplier_name"),
               "po_description": h.get("material_description") if len(its) == 1 else line.get("po_description"),
               "po_start": min(starts) if starts else None, "po_end": max(ends) if ends else None,
               "material_description": h.get("material_description"), "currency": h.get("currency"),
               "po_amount_doc": round(sum(_num(i.get("net_order_value")) for i in its), 2),
               "po_amount": round(sum(_num(i.get("value_inr")) for i in its), 2), "pr_no": h.get("pr_no")}
        if len(its) == 1:
            out.update(qty=h.get("quantity"), rate=h.get("net_price"))
        return {k: v for k, v in out.items() if k in ZMM_KEYS and v not in (None, "")}

    @staticmethod
    def _history(links, cover, items) -> Dict[str, List[Dict[str, Any]]]:
        """Per line, its links newest first (by item period start / created on) — for reusing the last split."""
        out: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
        for ln in links:
            its = cover.get(ln["_key"], {}).get("items", [])
            when = max([str(items[k].get("period_start") or items[k].get("created_on") or "") for k in its] or [""])
            out[ln["line_id"]].append({"po": ln["po"], "items": its, "alloc_pct": ln.get("alloc_pct"), "when": when})
        for v in out.values():
            v.sort(key=lambda x: x["when"], reverse=True)
        return out

    async def line_forecast_fy(self) -> float:
        cfg = await self.get_config()
        F = "F" + cfg["plan_fy"][2:]
        ms = fy_months(cfg["plan_fy"])
        return round(sum(sum(_num(f.get(vkey(F, p))) for p in ms) for f in (await self.fields("opex_tracker")).values()), 2)

    # ------------------------------------------------------------------ scope of review
    async def in_scope_items(self) -> set:
        """Items linked (any line, any status) or triaged with a type other than Mapping not required."""
        items = await self.fields("po_items")
        links = await self.fields("po_links")
        triage = await self.fields("po_triage")
        by_po = defaultdict(list)
        for k, it in items.items():
            by_po[it.get("po")].append(k)
        out = set()
        for ln in links.values():
            out |= set(by_po.get(str(ln.get("po")), []))
        for k, t in triage.items():
            if t.get("type") and t.get("type") != "Mapping not required":
                out.add(k)
            elif t.get("type") == "Mapping not required":
                out.discard(k)
        for ln in links.values():  # a link always keeps its items in review, even if triage says otherwise
            out |= set(by_po.get(str(ln.get("po")), []))
        return out

    # ------------------------------------------------------------------ the pipeline
    async def run_zmm_pipeline(self, content: bytes, filename: str, source: str = "manual",
                               meta: Optional[Dict[str, Any]] = None, by: str = "system") -> Dict[str, Any]:
        cfg = await self.get_config()
        meta = dict(meta or {})
        sha = hashlib.sha256(content).hexdigest()
        run_id = f"ZMM-{datetime.now(timezone.utc).strftime('%Y%m%d%H%M%S%f')[:-3]}"
        run = {"run_id": run_id, "source": source, "file_name": filename, "sha256": sha, "received_at": meta.get("received_at") or now_iso(),
               "status": "running", "by": by, **{k: v for k, v in meta.items() if k in ("message_id", "sender", "subject")}}
        await self.bulk_set("zmm_runs", {run_id: run}, upsert=True, by=by)
        try:
            out = await self._pipeline(content, filename, run_id, cfg, by)
            run.update(out, status="processed", processed_at=now_iso())
        except Exception as e:  # noqa: BLE001 — every failure is logged on the run and notified
            logger.exception("ZMM run %s failed", run_id)
            run.update(status="failed", error=str(e)[:2000], processed_at=now_iso())
        if self.storage:
            try:
                key = f"zmm/{run_id}/{filename or 'zmm.xlsx'}"
                await self.storage.save(key, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                run["storage_key"] = key
            except Exception as e:  # noqa: BLE001
                run["storage_error"] = str(e)[:300]
        await self.bulk_set("zmm_runs", {run_id: run}, by=by)
        return run

    async def _pipeline(self, content: bytes, filename: str, run_id: str, cfg: Dict[str, Any], by: str) -> Dict[str, Any]:
        plan = cfg["plan_fy"]
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        name = next((s for s in wb.sheetnames if s.strip().lower().replace(" ", "_") == "zmm_po_report"), wb.sheetnames[0])
        sheet_rows = [tuple(r) for r in wb[name].iter_rows(values_only=True, max_col=90)]
        header, raw = P.read_zmm(sheet_rows)
        if not raw:
            raise ValueError("The ZMM report has no PO rows")
        snapshot = max((str(f.get("created_on") or "") for f in raw), default="")
        before = await self.line_forecast_fy()
        P.register_rows(raw)
        P.final_values(raw)
        # 2. replace po_register
        await self.db.aop_rows.delete_many({"dataset": "po_register"})
        docs = [{"dataset": "po_register", "key": f.pop("_key"), "seq": i, "fields": f, "updated_at": now_iso(),
                 "updated_by": by} for i, f in enumerate(raw, start=1)]
        for j in range(0, len(docs), 1000):
            await self.db.aop_rows.insert_many(docs[j:j + 1000])
        # 4. po_items: latest values; accepted kept for items in review
        fx = await self.fx_table([d["fields"] for d in docs], by)
        fresh = P.build_items([d["fields"] for d in docs], fx.rate, plan)
        prev_items = await self.fields("po_items")
        first_run = not prev_items
        scope = await self.in_scope_items()
        mode = cfg.get("po_change_mode") or "hold"
        triage = await self.fields("po_triage")
        item_upd: Dict[str, Dict[str, Any]] = {}
        added_items, deleted_items = [], []
        for k, it in fresh.items():
            prev = prev_items.get(k)
            item_upd[k] = P.merge_item(it, prev, k in scope, mode)
            if prev is None and not first_run:
                added_items.append(k)
            if prev is not None and prev.get("latest_active", True) and not it["active"]:
                deleted_items.append(k)
        gone = []
        for k, prev in prev_items.items():
            if k in fresh:
                continue
            keep = k in scope or k in triage
            if keep:
                item_upd[k] = P.missing_item(prev, k in scope, mode)
            else:
                gone.append(k)
        if gone:
            await self.db.aop_rows.delete_many({"dataset": "po_items", "key": {"$in": gone}})
        await self.bulk_set("po_items", item_upd, upsert=True, by=by)
        # 3. triage enrichment + auto triage
        new_pos = await self._triage(docs, fresh, triage, cfg, by)
        # 5. PO changes on mapped items
        flagged, fx_only = await self._detect_changes(run_id, item_upd, scope, added_items, by)
        await self._corrections_after_run(item_upd, by)
        # 6–7. links + forecast
        res = await self.resolve_all()
        after = res.get("forecast_total", 0.0)
        await self.store_change_impacts()
        # 8. derived columns back onto po_register
        await self._write_back(docs)
        to_map = await self.count_to_map()
        return {"rows": len(raw), "po_items": len(fresh), "pos": len({f["purchase_order"] for f in raw}),
                "new_pos": new_pos, "items_added": len(added_items), "items_deleted": len(deleted_items),
                "items_missing": len([k for k, v in item_upd.items() if v.get("in_latest_zmm") is False]),
                "changes_flagged": flagged, "fx_only_changes": fx_only, "to_map": to_map, "snapshot": snapshot,
                "forecast_before": before, "forecast_after": after, "forecast_delta": round(after - before, 2),
                "first_run": first_run, "header": header[:120],
                "unconverted": len([1 for it in fresh.values() if P.fx_flagged(it.get("fx_source"))]),
                "fx_fetched": fx.fetched, "fx_fetch_error": fx.error}

    async def _triage(self, docs, fresh, triage, cfg, by) -> int:
        prefix = cfg.get("po_nature_prefix") or P.DEFAULT_PREFIX
        links = await self.fields("po_links")
        linked_pos = {str(f.get("po")) for f in links.values()}
        enrich_by_item: Dict[str, Dict[str, Any]] = {}
        for d in docs:
            f = d["fields"]
            k = f"{f['purchase_order']}|{f.get('purchase_order_item') or ''}"
            e = enrich_by_item.setdefault(k, {})
            for src, dst in P.ENRICH.items():
                if f.get(src) not in (None, "") and dst not in e:
                    e[dst] = f[src]
            if f.get("po_mapping_source") and "mapping_note" not in e:
                e["mapping_note"] = f["po_mapping_source"]
        upd: Dict[str, Dict[str, Any]] = {}
        new_pos = set()
        for k, it in fresh.items():
            cur = triage.get(k) or {}
            e = enrich_by_item.get(k, {})
            u = {kk: v for kk, v in e.items() if v not in (None, "") and cur.get(kk) != v}
            if not cur:
                new_pos.add(it["po"])
                u["created_at"] = now_iso()
                u["po"], u["item"] = it["po"], it["item"]
            if not cur.get("type"):
                if it["po"] in linked_pos:
                    u.update(type="Opex", opex_action="Renewal / replacement", decided_by="mapping", decided_at=now_iso())
                elif not it["active"]:
                    u.update(type="Mapping not required", not_required_reason="Deleted / blocked", decided_by="system",
                             decided_at=now_iso())
                else:
                    pre = P.type_from_wbs(it.get("wbs"), prefix, e.get("nature") or cur.get("nature"))
                    if pre and cur.get("suggested_type") != pre:
                        u["suggested_type"] = pre
            if u:
                upd[k] = u
        await self.bulk_set("po_triage", upd, upsert=True, by=by)
        return len(new_pos - {t.get("po") for t in triage.values()})

    async def count_to_map(self) -> int:
        n = 0
        items = await self.fields("po_items")
        for k, t in (await self.fields("po_triage")).items():
            it = items.get(k)
            if it and it.get("active", True) and it.get("in_latest_zmm", True) is not False and not t.get("type"):
                n += 1
        return n

    async def _detect_changes(self, run_id, item_upd, scope, added, by) -> Tuple[int, int]:
        open_ch = {(c.get("po"), c.get("item"), c.get("field")): (k, c)
                   for k, c in (await self.fields("po_changes", {"fields.status": "Open"})).items()}
        rejected = {(c.get("po"), c.get("item"), c.get("field")): c
                    for c in (await self.fields("po_changes", {"fields.status": "Rejected"})).values()}
        links = await self.fields("po_links")
        triage = await self.fields("po_triage")
        linked_pos = defaultdict(set)
        for f in links.values():
            linked_pos[str(f.get("po"))].add(f.get("line_id"))
        upd: Dict[str, Dict[str, Any]] = {}
        closed: Dict[str, Dict[str, Any]] = {}
        flagged = fx_only = 0
        seen = set()
        held = {(po, item) for (po, item, fld) in open_ch if fld == "item_added"}
        for k, it in item_upd.items():
            if k not in scope or tuple(k.split("|", 1)) in held:
                continue
            po, item = k.split("|", 1)
            for ch in P.diff_item(it):
                key3 = (po, item, ch["field"])
                seen.add(key3)
                rj = rejected.get(key3)
                if rj and rj.get("new") == ch["new"]:
                    continue  # already rejected this exact value
                if key3 in open_ch:
                    ck, c = open_ch[key3]
                    if c.get("new") != ch["new"]:
                        upd[ck] = {"new": ch["new"], "old": ch["old"], "updated_run": run_id}
                else:
                    ck = f"{run_id}|{po}|{item}|{ch['field']}"
                    upd[ck] = {"run_id": run_id, "po": po, "item": item, "field": ch["field"],
                               "label": P.FIELD_LABEL.get(ch["field"], ch["field"]), "old": ch["old"], "new": ch["new"],
                               "status": "Open", "lines": sorted(linked_pos.get(po, set())), "created_at": now_iso()}
                    flagged += 1
            if P.fx_only_change(it) and not P.diff_item(it):
                fx_only += 1
                upd_fx = P.accept_fields(it, ["value_inr", "fx_source"])
                item_upd[k].update(upd_fx)
                await self.bulk_set("po_items", {k: upd_fx}, by=by)
        # items added to a mapped PO
        for k in added:
            po, item = k.split("|", 1)
            if po not in linked_pos and not any(t.get("po") == po and t.get("type") for t in triage.values()):
                continue
            key3 = (po, item, "item_added")
            if key3 in open_ch:
                continue
            prior = next((t for t in triage.values() if t.get("po") == po and t.get("type")), {})
            ck = f"{run_id}|{po}|{item}|item_added"
            it = item_upd[k]
            upd[ck] = {"run_id": run_id, "po": po, "item": item, "field": "item_added", "label": P.FIELD_LABEL["item_added"],
                       "old": None, "new": f"{it.get('material') or ''} {it.get('material_description') or ''} ₹{_num(it.get('latest_value_inr')):,.0f}".strip(),
                       "status": "Open", "lines": sorted(linked_pos.get(po, set())), "created_at": now_iso(),
                       "prefill": {"type": prior.get("type") or ("Opex" if po in linked_pos else None),
                                   "line_id": sorted(linked_pos.get(po, set()))[0] if linked_pos.get(po) else prior.get("line_id")}}
            flagged += 1
            # a new item on a mapped PO is held out of the forecast until accepted
            hold = {"accepted_active": False}
            item_upd[k].update(hold)
            await self.bulk_set("po_items", {k: hold}, by=by)
        # open changes whose latest value returned to the accepted one close themselves
        for key3, (ck, c) in open_ch.items():
            if key3[2] == "item_added" or key3 in seen or ck in upd:
                continue
            closed[ck] = {"status": "Reverted", "closed_at": now_iso(), "decided_by": "SAP (value reverted)"}
        await self.bulk_set("po_changes", upd, upsert=True, by=by)
        await self.bulk_set("po_changes", closed, by=by)
        return flagged, fx_only

    async def _corrections_after_run(self, item_upd, by):
        """A ZMM run sets 'Possibly resolved' when the field named by the correction type changed."""
        upd = {}
        for k, c in (await self.fields("po_corrections")).items():
            if c.get("status") not in ("Open", "Raised"):
                continue
            fields = P.CORRECTION_FIELDS.get(c.get("type") or "", [])
            snap = c.get("snapshot") or {}
            for it in c.get("items") or []:
                cur = item_upd.get(f"{c.get('po')}|{it}") or {}
                if any(f in snap and not P._same(f, snap.get(f), cur.get(f"latest_{f}")) for f in fields):
                    upd[k] = {"status": "Possibly resolved", "possibly_resolved_at": now_iso()}
                    break
        await self.bulk_set("po_corrections", upd, by=by)

    async def _write_back(self, docs):
        """Linked Old PO, Linked Forecast S.No, Link Status, Amount Match, Active (excl L/S) on po_register."""
        links = await self.fields("po_links")
        lines = await self.fields("opex_tracker")
        items = await self.fields("po_items")
        by_po = defaultdict(list)
        for ln in links.values():
            by_po[str(ln.get("po"))].append(ln)
        upd = {}
        for d in docs:
            f = d["fields"]
            po = f["purchase_order"]
            lns = [ln for ln in by_po.get(po, []) if not ln.get("material") or ln.get("material") == f.get("material")]
            it = items.get(f"{po}|{f.get('purchase_order_item') or ''}") or {}
            old = sorted({str((lines.get(ln["line_id"]) or {}).get("po") or "") for ln in lns} - {""})
            alloc = sum(_num(ln.get("alloc_value_inr")) for ln in lns)
            val = sum(_num(ln.get("value_inr")) for ln in lns)
            upd[d["key"]] = {"linked_old_po": ", ".join(old) or None,
                             "linked_forecast_s_no": ", ".join(sorted({ln["line_id"] for ln in lns})) or None,
                             "link_status": ("Linked" if lns else "Not linked"),
                             "amount_match": (None if not lns else ("Matched" if val and abs(alloc - val) <= max(1.0, 0.005 * val) else "Partly allocated")),
                             "active_excl_l_s": "Yes" if it.get("active", True) else "No"}
        await self.bulk_set("po_register", upd, by="system")


# ---------------------------------------------------------------------------------------------- grid columns
def _c(key, label, ctype="text", editable=False, hidden=False, role=None):
    c = column(key, label, ctype, editable=editable, hidden=hidden)
    c["role"] = role or ("input" if editable else "computed")
    return c


PO_META: Dict[str, List[Dict[str, Any]]] = {
    "fx_rates": [_c("currency", "Currency", role="key"), _c("date", "Date", "date", role="key"),
                 _c("rate", "INR per unit", "number", True), _c("source", "Source", editable=True)],
    "po_links": [
        _c("line_id", "Line ID", role="key"), _c("po", "PO", role="key"), _c("material", "Material code", role="key"),
        _c("po_item", "PO item", role="key"), _c("alloc_pct", "Allocation %", "percent", True),
        _c("alloc_auto", "Auto-split"), _c("coverage_from", "Coverage from", "date", True),
        _c("coverage_to", "Coverage to", "date", True), _c("status", "Link status", editable=True),
        _c("remarks", "Remarks", editable=True), _c("items", "Items covered"), _c("value_inr", "PO value (INR)", "number"),
        _c("alloc_value_inr", "Allocated (INR)", "number"), _c("period_start", "Period start", "date"),
        _c("period_end", "Period end", "date"), _c("supplier_code", "Supplier code"), _c("supplier_name", "Supplier"),
        _c("supplier_changed", "Supplier changed"), _c("grn_inr", "GRN (INR)", "number"),
        _c("pending_inr", "Pending (INR)", "number"), _c("fy_impact_inr", "FY impact (INR)", "number"),
        _c("flags", "Flags"), _c("source", "Source"), _c("resolved_at", "Resolved at", hidden=True)],
    "po_items": [
        _c("po", "PO", role="key"), _c("item", "Item", role="key"), _c("material", "Material"),
        _c("material_description", "Material description"), _c("supplier_code", "Supplier code"),
        _c("supplier_name", "Supplier"), _c("created_on", "Created on", "date"), _c("pr_no", "PR"), _c("wbs", "WBS"),
        _c("aop_code", "AOP code (Short ID)"), _c("wbs_name", "WBS name"), _c("material_type", "Material type"),
        _c("currency", "Currency"), _c("quantity", "Qty", "number"), _c("net_price", "Net order price", "number"),
        _c("net_order_value", "Net order value", "number"),
        _c("value_inr", "Value (INR)", "number"), _c("fx_rate", "FX rate (PO date)", "number"), _c("fx_date", "FX at (PO date)", "date"), _c("fx_source", "FX source"), _c("period_start", "Period start", "date"),
        _c("period_end", "Period end", "date"), _c("delivery_date", "Delivery date", "date"), _c("active", "Active"),
        _c("deletion_indicator", "Deletion indicator"), _c("grn_inr", "GRN (INR)", "number"),
        _c("pending_inr", "Pending (INR)", "number"), _c("grn_pct", "GRN %", "percent"),
        _c("invoiced_inr", "Invoiced (INR)", "number"), _c("last_grn_date", "Last GRN", "date"),
        _c("last_invoice_date", "Last invoice", "date"), _c("fy_impact_inr", "FY impact (INR)", "number"),
        _c("in_latest_zmm", "In latest ZMM"), _c("accepted_value_inr", "Accepted value (INR)", "number", hidden=True),
        _c("accepted_period_start", "Accepted start", "date", hidden=True),
        _c("accepted_period_end", "Accepted end", "date", hidden=True)],
    "po_triage": [
        _c("po", "PO", role="key"), _c("item", "Item", role="key"), _c("type", "Type", editable=True),
        _c("suggested_type", "Suggested type"), _c("opex_action", "Opex action", editable=True),
        _c("line_id", "Line", editable=True), _c("department", "Department", editable=True),
        _c("budget_code", "Budget code", editable=True), _c("capex_key", "Capex project", editable=True),
        _c("not_required_reason", "Not required reason", editable=True), _c("needs_correction", "Needs correction"),
        _c("category", "Category"), _c("location", "Location"), _c("nature", "Nature (SAP file)"),
        _c("nature_of_services", "Nature of services"), _c("mapping_note", "PO mapping (source)"),
        _c("remarks", "Remarks", editable=True), _c("decided_by", "Decided by"), _c("decided_at", "Decided at")],
    "po_changes": [
        _c("po", "PO", role="key"), _c("item", "Item", role="key"), _c("label", "Change"), _c("old", "Accepted"),
        _c("new", "SAP now"), _c("status", "Status"), _c("lines", "Lines"), _c("fy_impact", "FY impact", "number"),
        _c("remarks", "Remarks"), _c("run_id", "Run", role="key"), _c("field", "Field", hidden=True, role="key"),
        _c("decided_by", "Decided by"), _c("decided_at", "Decided at")],
    "po_corrections": [
        _c("line_id", "Correction ID", role="key"), _c("po", "PO"), _c("items", "Items"), _c("supplier", "Supplier"),
        _c("type", "Correction type", editable=True), _c("remarks", "Remarks", editable=True),
        _c("raised_to", "Raised to", editable=True), _c("raised_on", "Raised on", "date", True),
        _c("status", "Status", editable=True), _c("override", "Provisional override"), _c("source", "Source"),
        _c("resolved_on", "Resolved on", "date"), _c("resolved_by", "Resolved by")],
    "zmm_runs": [
        _c("run_id", "Run", role="key"), _c("source", "Source"), _c("received_at", "Received at"),
        _c("processed_at", "Processed at"), _c("file_name", "File"), _c("status", "Status"), _c("error", "Error"),
        _c("rows", "Rows", "number"), _c("po_items", "PO items", "number"), _c("new_pos", "New POs", "number"),
        _c("items_added", "Items added", "number"), _c("items_deleted", "Items deleted", "number"),
        _c("changes_flagged", "Changes flagged", "number"), _c("to_map", "To map", "number"),
        _c("forecast_delta", "Forecast Δ (INR)", "number"), _c("sha256", "SHA-256", hidden=True)],
}


# ---------------------------------------------------------------------------------------------- notifications
def run_summary(run: Dict[str, Any]) -> Tuple[str, str]:
    if run.get("status") == "failed":
        return "ZMM run failed", f"{run.get('file_name')}: {run.get('error')}"
    d = _num(run.get("forecast_delta"))
    return ("ZMM processed",
            f"{run.get('to_map') or 0} new PO items to map, {run.get('changes_flagged') or 0} PO changes to review, "
            f"forecast Δ ₹{d / 1e5:,.2f} L")


async def notify_zmm(db, run: Dict[str, Any]):
    """E-mail the admins (Graph mailer) and drop an in-app notification for each, linking to Review."""
    import uuid
    title, body = run_summary(run)
    admins = await db.users.find({"role": "admin"}, {"_id": 0, "id": 1, "email": 1}).to_list(200)
    docs = [{"id": str(uuid.uuid4()), "user_id": a["id"], "kind": "zmm_run", "title": title, "body": body,
             "link": "/app/aop/review", "read": False, "created_at": now_iso()} for a in admins if a.get("id")]
    if docs:
        await db.notifications_inapp.insert_many(docs)
    try:
        from notifications import mailer
        html = (f"<p><b>{title}</b></p><p>{body}</p><p>Source: {run.get('source')} · file {run.get('file_name')} · "
                f"run {run.get('run_id')}</p><p>Open FinSight → Annual Operating Plan → Review.</p>")
        res = await mailer.send(f"FinSight — {title}", html, [a["email"] for a in admins if a.get("email")])
        return res
    except Exception as e:  # noqa: BLE001 — a mail failure never fails the run
        logger.warning("ZMM notification mail failed: %s", e)
        return {"sent": False, "error": str(e)}
